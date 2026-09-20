"""Text extraction for PDF / DOCX / images, with OCR fallback for scanned pages."""

import logging
from dataclasses import dataclass, field
from pathlib import Path

from app.extraction.ocr import get_ocr_engine
from app.ingestion.validators import DOCX_MIME

log = logging.getLogger(__name__)

# A PDF page with fewer extractable characters than this is treated as scanned.
MIN_TEXT_CHARS_PER_PAGE = 40
OCR_RENDER_DPI = 200


@dataclass
class PageText:
    number: int  # 1-based
    text: str
    ocr: bool = False
    ocr_confidence: float | None = None


@dataclass
class ExtractedText:
    pages: list[PageText] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(p.text for p in self.pages)

    @property
    def ocr_used(self) -> bool:
        return any(p.ocr for p in self.pages)

    @property
    def page_count(self) -> int:
        return len(self.pages)

    def page_offsets(self) -> list[tuple[int, int, int]]:
        """(page_number, start, end) char spans within `self.text`."""
        spans, pos = [], 0
        for p in self.pages:
            spans.append((p.number, pos, pos + len(p.text)))
            pos += len(p.text) + 2  # the "\n\n" join
        return spans


def extract_text(path: Path, mime_type: str, *, ocr_enabled: bool = True) -> ExtractedText:
    if mime_type == "application/pdf":
        return _extract_pdf(path, ocr_enabled=ocr_enabled)
    if mime_type == DOCX_MIME:
        return _extract_docx(path)
    if mime_type.startswith("image/"):
        return _extract_image(path)
    raise ValueError(f"Unsupported MIME type for extraction: {mime_type}")


def _extract_pdf(path: Path, *, ocr_enabled: bool) -> ExtractedText:
    import pymupdf as fitz

    result = ExtractedText()
    with fitz.open(path) as doc:
        for index, page in enumerate(doc, start=1):
            text = page.get_text("text") or ""
            if len(text.strip()) >= MIN_TEXT_CHARS_PER_PAGE or not ocr_enabled:
                result.pages.append(PageText(number=index, text=text))
                continue
            log.info("Page %s of %s has little text; running OCR", index, path.name)
            pix = page.get_pixmap(dpi=OCR_RENDER_DPI)
            ocr = get_ocr_engine().read(pix.tobytes("png"))
            result.pages.append(
                PageText(number=index, text=ocr.text, ocr=True, ocr_confidence=ocr.confidence)
            )
    return result


def _extract_docx(path: Path) -> ExtractedText:
    import docx

    document = docx.Document(str(path))
    lines: list[str] = []
    for para in document.paragraphs:
        if para.text.strip():
            lines.append(para.text.strip())
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            # Drop merged-cell duplicates python-docx reports for spanning cells.
            deduped = [c for i, c in enumerate(cells) if i == 0 or c != cells[i - 1]]
            deduped = [c for c in deduped if c]
            if not deduped:
                continue
            # Two-column rows are almost always "Label | Value" in term sheets.
            lines.append(": ".join(deduped) if len(deduped) == 2 else " | ".join(deduped))
    return ExtractedText(pages=[PageText(number=1, text="\n".join(lines))])


def _extract_image(path: Path) -> ExtractedText:
    ocr = get_ocr_engine().read(path.read_bytes())
    return ExtractedText(
        pages=[PageText(number=1, text=ocr.text, ocr=True, ocr_confidence=ocr.confidence)]
    )
